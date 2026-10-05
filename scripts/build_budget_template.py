#!/usr/bin/env python3
"""Шаблон «Таблица доходов и расходов» для /tablica-dohodov-i-rashodov/.

Собирает landing/files/tablica-dohodov-i-rashodov-kubysh.xlsx. Логика та же, что в Кубыше:
период от получки до получки, доходы − обязательные − цели = деньги на жизнь, ÷ дни = бюджет на день.
Запуск: python3 scripts/build_budget_template.py (нужен openpyxl).
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).resolve().parents[1] / "landing" / "files" / "tablica-dohodov-i-rashodov-kubysh.xlsx"
LIME = PatternFill("solid", fgColor="CEF16C")
GREY = PatternFill("solid", fgColor="F2F2F6")
INPUT = PatternFill("solid", fgColor="FFFBE6")
BOLD = Font(bold=True)
TITLE = Font(bold=True, size=16)
THIN = Border(bottom=Side(style="thin", color="DDDDDD"))
RUB = '# ##0 "₽"'
CATEGORIES = ["Продукты", "Кафе и доставка", "Транспорт", "Дом", "Здоровье", "Одежда", "Развлечения", "Подарки", "Другое"]

wb = Workbook()
ws = wb.active
ws.title = "Бюджет периода"
ws.column_dimensions["A"].width = 38
ws.column_dimensions["B"].width = 16
ws.column_dimensions["C"].width = 16
ws.column_dimensions["D"].width = 44

ws["A1"] = "Таблица доходов и расходов от получки до получки"
ws["A1"].font = TITLE
ws["A2"] = "Желтые ячейки заполняете вы, остальное считается само. Шаблон от kubysh.com"
ws["A2"].font = Font(italic=True, color="666666")

row = 4


def block(title, items, note=""):
    """Таблица раздела: название, сумма, комментарий. Возвращает адрес ячейки с итогом."""
    global row
    ws.cell(row, 1, title).font = BOLD
    ws.cell(row, 2, "Сумма").font = BOLD
    ws.cell(row, 4, note).font = Font(italic=True, color="666666")
    for c in range(1, 5):
        ws.cell(row, c).fill = LIME
    row += 1
    first = row
    for name, value, comment in items:
        ws.cell(row, 1, name)
        cell = ws.cell(row, 2, value)
        cell.number_format = RUB
        cell.fill = INPUT
        ws.cell(row, 4, comment).font = Font(color="666666")
        for c in range(1, 5):
            ws.cell(row, c).border = THIN
        row += 1
    ws.cell(row, 1, "Итого").font = BOLD
    total = ws.cell(row, 2, f"=SUM(B{first}:B{row - 1})")
    total.number_format = RUB
    total.font = BOLD
    for c in range(1, 5):
        ws.cell(row, c).fill = GREY
    ref = f"B{row}"
    row += 2
    return ref


income = block("1. Доходы за период", [
    ("Аванс", 32000, "10-го числа"),
    ("Зарплата", 48000, "25-го числа"),
    ("Премия, подработка", 0, "если есть в этом периоде"),
    ("Другое", 0, ""),
], "Считайте период от одной получки до следующей")
mandatory = block("2. Обязательные платежи", [
    ("Аренда или ипотека", 25000, ""),
    ("Коммунальные услуги", 0, "если платите отдельно"),
    ("Связь и интернет", 1500, ""),
    ("Кредиты и рассрочки", 6000, ""),
    ("Подписки", 1000, "музыка, кино, облако"),
    ("Транспорт: проездной", 1500, ""),
], "То, что уйдет в любом случае")
goals = block("3. Отложить на цели", [
    ("Финансовая подушка", 9000, "первая по приоритету"),
    ("Цель: отпуск", 6000, ""),
    ("Цель: другое", 0, ""),
], "Откладывайте в день получки, а не в конце месяца")

ws.cell(row, 1, "4. Деньги на жизнь до следующей получки").font = BOLD
free = ws.cell(row, 2, f"={income}-{mandatory}-{goals}")
free.number_format = RUB
free.font = Font(bold=True, size=13)
for c in range(1, 5):
    ws.cell(row, c).fill = LIME
free_ref = f"B{row}"
row += 1
ws.cell(row, 1, "Дней до следующей получки")
days = ws.cell(row, 2, 30)
days.fill = INPUT
days_ref = f"B{row}"
row += 1
ws.cell(row, 1, "Бюджет на день").font = BOLD
per_day = ws.cell(row, 2, f"=IF({days_ref}>0,{free_ref}/{days_ref},0)")
per_day.number_format = RUB
per_day.font = Font(bold=True, size=13)
ws.cell(row, 4, "Столько можно тратить каждый день, чтобы хватило до получки").font = Font(color="666666")
row += 2

ws.cell(row, 1, "5. Факт по тратам").font = BOLD
for c in range(1, 5):
    ws.cell(row, c).fill = LIME
row += 1
ws.cell(row, 1, "Потрачено (лист «Траты»)")
spent = ws.cell(row, 2, "=SUM('Траты'!C2:C301)")
spent.number_format = RUB
spent_ref = f"B{row}"
row += 1
ws.cell(row, 1, "Осталось на жизнь").font = BOLD
left = ws.cell(row, 2, f"={free_ref}-{spent_ref}")
left.number_format = RUB
left.font = BOLD
row += 2

ws.cell(row, 1, "Траты по категориям").font = BOLD
ws.cell(row, 2, "Сумма").font = BOLD
ws.cell(row, 3, "Доля").font = BOLD
for c in range(1, 5):
    ws.cell(row, c).fill = GREY
row += 1
for cat in CATEGORIES:
    ws.cell(row, 1, cat)
    s = ws.cell(row, 2, f"=SUMIF('Траты'!B2:B301,A{row},'Траты'!C2:C301)")
    s.number_format = RUB
    share = ws.cell(row, 3, f"=IF({spent_ref}>0,B{row}/{spent_ref},0)")
    share.number_format = "0%"
    row += 1

row += 1
ws.cell(row, 1, "Не хочется заполнять руками? В приложении Кубыш траты вносятся скриншотом из банка, а бюджет на день пересчитывается сам: kubysh.com").font = Font(italic=True, color="3D6B00")

# Лист трат
tr = wb.create_sheet("Траты")
for col, (title, width) in enumerate([("Дата", 14), ("Категория", 22), ("Сумма", 14), ("Комментарий", 40)], start=1):
    cell = tr.cell(1, col, title)
    cell.font = BOLD
    cell.fill = LIME
    tr.column_dimensions[cell.column_letter].width = width
examples = [("10.10.2026", "Продукты", 2350, "магазин у дома"), ("10.10.2026", "Транспорт", 320, "такси"), ("11.10.2026", "Кафе и доставка", 890, "обед")]
for r, (d, cat, amount, comment) in enumerate(examples, start=2):
    tr.cell(r, 1, d)
    tr.cell(r, 2, cat)
    tr.cell(r, 3, amount).number_format = RUB
    tr.cell(r, 4, comment)
for r in range(2, 302):
    tr.cell(r, 3).number_format = RUB
dv = DataValidation(type="list", formula1='"' + ",".join(CATEGORIES) + '"', allow_blank=True)
tr.add_data_validation(dv)
dv.add("B2:B301")
tr.freeze_panes = "A2"

# Инструкция
how = wb.create_sheet("Как пользоваться")
how.column_dimensions["A"].width = 100
lines = [
    ("Как пользоваться таблицей", TITLE),
    ("1. Период считайте от получки до получки, а не по календарному месяцу. Если аванс и зарплата приходят в разные дни, ведите один лист на период между ними или на весь месяц, как удобнее.", None),
    ("2. На листе «Бюджет периода» впишите доходы, обязательные платежи и суммы на цели в желтые ячейки.", None),
    ("3. Укажите, сколько дней до следующей получки. Таблица посчитает деньги на жизнь и бюджет на день.", None),
    ("4. Каждую трату записывайте на листе «Траты»: дата, категория из списка, сумма. Остаток и доли по категориям посчитаются сами.", None),
    ("5. В новый период скопируйте лист или файл и обнулите траты.", None),
    ("Google Таблицы: откройте sheets.google.com → Файл → Импорт → Загрузка → выберите этот файл.", None),
    ("Шаблон сделан командой приложения Кубыш (kubysh.com). Кубыш делает то же самое в iPhone автоматически: траты скриншотом из банка, бюджет на день, цели с датами.", None),
]
for i, (text, font) in enumerate(lines, start=1):
    c = how.cell(i, 1, text)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    if font:
        c.font = font

OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print(OUT, OUT.stat().st_size)
